import sys
sys.path.insert(0, '/private/tmp/geldium_pptx_lib')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

OUT = Path(__file__).parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY='102839'; TEAL='007F80'; WHITE='FFFFFF'; INK='233E4D'; MUTED='55707C'; PALE='EAF2F4'; GOLD='E2B454'
def rect(s,x,y,w,h,color):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor.from_string(color); a.line.fill.background()
    return a
def txt(s,x,y,w,h,text,size=18,color=INK,bold=False):
    box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.03); tf.margin_top=Inches(.02)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(10)
    return box
def slide(n,title,sub,source):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    rect(s,0,0,13.333,.14,TEAL)
    txt(s,.55,.35,12,.3,'GELDIUM  /  RESPONSIBLE COLLECTIONS',11,TEAL,True)
    txt(s,.55,.86,12.2,.62,title,30,NAVY,True)
    txt(s,.55,1.58,12.1,.62,sub,16,MUTED)
    rect(s,.55,6.94,12.2,.015,'CCDADD')
    txt(s,.55,7.04,11.8,.27,source,10,MUTED)
    txt(s,12.35,7.02,.45,.3,f'{n:02}',12,TEAL,True)
    return s
def note(s,text): s.notes_slide.notes_text_frame.text=text
def banner(s,text,y=6.12):
    rect(s,.55,y,12.2,.63,NAVY); txt(s,.77,y+.11,11.75,.43,text,16,WHITE,True)

s=slide(1,'A controlled loop turns risk signals into support','Executive concept: prevent avoidable delinquency through early, proportionate customer assistance.','Sources: EDA_SummaryReport.docx; Task 2_ModelPlan.docx | Proposed workflow')
cards=[('01  CUSTOMER DATA','Payment history, balances, income and utilization.\nAdd contact preferences and verified account status.\nValidate quality and timing.'),('02  DECISION','Interpretable logistic regression estimates risk.\nPolicy checks select an eligible action.\nUncertain cases go to staff.'),('03  ACTION','Send approved reminders through permitted channels.\nOffer a support conversation.\nRoute hardship or disputes to staff.'),('04  LEARNING','Track repayment, response, complaints and fairness.\nCompare with a control group.\nHumans approve model updates.')]
for i,(h,b) in enumerate(cards):
    x=.55+i*3.1; rect(s,x,2.5,2.9,3.22,PALE); rect(s,x,2.5,2.9,.08,TEAL)
    txt(s,x+.18,2.79,2.53,.5,h,17,TEAL,True); txt(s,x+.18,3.42,2.53,2.17,b,18)
    if i<3: txt(s,x+2.91,3.8,.2,.5,'›',24,TEAL,True)
banner(s,'Learning feeds the next decision cycle; policy and model changes always require approval.')
note(s,'Outline: slides 1–2 workflow and decision policy; slide 3 agent autonomy; slide 4 guardrails; slide 5 impact. The current extract contains 500 customers, of whom 80 (16%) are delinquent. This is descriptive evidence, not a validated forecast. The Task 2 model is a proposal, with no fitted coefficients or evaluation results supplied. Contact preferences and verified account status are proposed operational inputs, not fields confirmed in the extract. The predictive model scores risk; the agent executes approved workflows. Keep these responsibilities distinct. Sources: EDA_SummaryReport.docx; Task 2_ModelPlan.docx. Updated_Business_Summary_Report.docx was read but contains unfilled template prompts, so there are no completed Task 3 recommendations to claim as established.')

s=slide(2,'Risk prioritizes help; policy determines the action','Resolve data limitations first, then test customer support pathways in a controlled pilot.','Sources: EDA_SummaryReport.docx; Task 2_ModelPlan.docx | Proposed pathways; thresholds to be validated')
rect(s,.55,2.38,4.1,3.48,PALE)
txt(s,.77,2.62,3.65,.45,'BEFORE AUTOMATION',18,TEAL,True)
txt(s,.77,3.18,3.65,2.52,'• Repair On-time / Late / Missed encoding.\n• Standardize employment labels; preserve missingness flags.\n• Verify history precedes the outcome.\n• Validate and calibrate risk scores; set thresholds to staff capacity.',18)
txt(s,4.95,2.43,7.55,.38,'ELIGIBILITY → RISK → SUPPORT',17,TEAL,True)
rows=[('Low validated risk','Routine due-date reminder; respect channel preferences.'),('Elevated validated risk','Proactive support invitation; staff review repayment options.'),('Missing balance / uncertainty','Verify source data before using the risk score to act.'),('Hardship, dispute or restrictions','Pause automated collection outreach; specialist review.')]
for i,(a,b) in enumerate(rows):
    y=3.02+i*.7; rect(s,4.95,y,7.8,.64,PALE if i%2==0 else 'F6F8F9')
    txt(s,5.08,y+.06,2.6,.54,a,15,INK,True); txt(s,7.83,y+.06,4.76,.54,b,15)
banner(s,'Observed signal: missing loan balance = 24.1% delinquency vs. 16.0% overall; verify, then assist.')
note(s,'EDA associations are exploratory: all numeric pairwise correlations have |rho| <= 0.047. Missing loan balance shows 24.1% delinquency (1.51x overall), Business card holders 21.3%, and tenure 5–10 years 19.7%. Use these as investigation hypotheses, not causal effects or validated individual-risk rules. Do not prioritize collection pressure based on age. Start with the Task 2 L2-regularized logistic regression and class weighting. Keep raw missing data visible; median imputation for fitting and scaling must occur inside training folds, with missingness indicators. Use stratified five-fold validation, precision, recall, F1 and ROC-AUC; assess calibration because class weighting can affect probability estimates. Add a later-period evaluation once dated data are available. Define outcome horizon and exclude future information. No numerical thresholds are justified by the supplied documents. The pilot intervention—supportive reminders and staff-led repayment support—is a new proposal because Task 3 is unfilled. Suppression eligibility comes before scoring. Missingness may be informative but may also reflect operational bias; it must not automatically trigger more aggressive contact.')

s=slide(3,'Agents execute approved steps; people own exceptions','Bounded autonomy increases scale while retaining accountability for consequential decisions.','Source: Task 2_ModelPlan.docx | Proposed operating model')
for x,h,c in [(.55,'AUTONOMOUS',TEAL),(6.75,'HUMAN OVERSIGHT',NAVY)]:
    rect(s,x,2.45,6.02,.65,c); txt(s,x+.2,2.59,5.61,.43,h,20,WHITE,True)
left=['Validate inputs; score eligible accounts with the approved model.','Select approved channel, timing and reminder template.','Send reminders; record responses and create staff tasks.','Log decisions; monitor drift, complaints and control breaches.']
right=['Approve model, thresholds, contact rules and message templates.','Review hardship, disputes, uncertain scores and data anomalies.','Authorize plan changes, settlements and any adverse or legal action.','Approve retraining and rollout; own appeals, audits and shutdown.']
for i in range(4):
    y=3.13+i*.65
    for x,t in [(.55,left[i]),(6.75,right[i])]:
        rect(s,x,y,6.02,.62,PALE if i%2==0 else 'F6F8F9'); txt(s,x+.18,y+.06,5.66,.54,t,16)
banner(s,'Exception → pause outreach → queue a named specialist → record the decision and customer explanation.')
note(s,'Agentic AI plans and carries out a sequence of allowed activities within a fixed policy: eligibility checks, approved reminder scheduling, response tracking and task creation. It cannot invent balances, fees, legal claims, discounts or promises. Customer messages are untrusted inputs and cannot override policy or tool permissions. Use a narrow tool allowlist, approved message templates, verified account facts and action logs. Escalate requests outside the approved menu, vulnerable-customer indicators, disputed debt, conflicting data and repeated nonresponse. Compliance and Collections set review service levels before launch; the operations lead owns the queue and can disable automation. Model Risk owns validation and release approval. Human reviewers need the data, risk reasons and action history, and must be able to correct or override the recommendation. Retraining proposals can be prepared automatically; no live model changes without approval.')

s=slide(4,'Four guardrails make automation accountable','Embed safeguards in each decision and maintain a clear route to human review.','References: NIST AI RMF; CFPB Debt Collection Rule FAQs | URLs and applicability details in speaker notes')
guards=[('01  FAIRNESS','Audit false-positive rates and support access across age, location and employment groups. Investigate proxies; pause affected automation when approved limits are breached.'),('02  EXPLAINABILITY','Provide plain-language risk reasons and a route to correct data or appeal. Log input facts, model version, action, reason and any human override.'),('03  COMPLIANCE & PRIVACY','Compliance approves rules for each market and creditor role: permitted contact, frequency, opt-outs and disputes. Minimize data; restrict access and retention.'),('04  CONTROL & ACCOUNTABILITY','Use verified account facts and approved templates. Route hardship and adverse actions to staff; monitor incidents and provide an immediate shutdown switch.')]
for i,(h,b) in enumerate(guards):
    y=2.37+i*.87; rect(s,.55,y,12.2,.78,PALE)
    txt(s,.76,y+.11,3.12,.55,h,17,TEAL,True); txt(s,3.98,y+.09,8.52,.65,b,17)
banner(s,'Scale only after Model Risk, Compliance and Collections approve the evidence and operating controls.')
note(s,'Fairness reviews should use subgroup sample counts and uncertainty: n=500 can produce unstable differences, so do not claim demonstrated parity. Audit age, location and employment status and demographic-correlated inputs; do not use age to target collection pressure. Missingness indicators also need bias review. Agree breach thresholds before launch; review weekly during the pilot. Protect demographic audit data with restricted access. Explainability should reflect actual contributing inputs, not unsupported causal claims. NIST AI RMF is voluntary risk-management guidance, not a certification or legal compliance guarantee: https://www.nist.gov/itl/ai-risk-management-framework . CFPB guidance supports attention to inconvenient contact times and places, electronic opt-outs and collection restrictions: https://www.consumerfinance.gov/compliance/compliance-resources/other-applicable-requirements/debt-collection/debt-collection-rule-faqs/ . The dataset lists US cities, but Geldium jurisdiction and creditor/collector role are not confirmed. Compliance must determine whether FDCPA/Regulation F, relevant state rules and other obligations apply before enabling actions. The slide describes proposed controls, not a legal determination. Do not assume a uniform legal regime or encode contact limits without review.')

s=slide(5,'Prove value before expanding automation','Proposed 90-day controlled pilot after validation; targets below are hypotheses, not forecasts.','Sources: EDA_SummaryReport.docx; Task 2_ModelPlan.docx | Illustrative targets; no measured uplift supplied')
for x,h,c in [(.55,'BUSINESS KPIs',TEAL),(6.75,'CUSTOMER OUTCOMES',NAVY)]:
    rect(s,x,2.39,6.02,.62,c); txt(s,x+.18,2.52,5.66,.44,h,20,WHITE,True)
txt(s,.77,3.18,5.55,2.48,'• 10% relative reduction in new delinquency versus concurrent control.\n• 15% lower cost per resolved account, including AI and review costs.\n• Track cure rate and net recovery per eligible account.\n• Reduce routine workload; scale within staff capacity.',19)
txt(s,6.97,3.18,5.55,2.48,'• Earlier access to affordable repayment support.\n• Clear messages, preferred channels and simple opt-outs.\n• Complaint rate no worse than control; improve customer satisfaction.\n• No material worsening in subgroup false-positive or support-access gaps.',19)
rect(s,.55,5.89,12.2,.84,PALE)
txt(s,.76,6.04,11.74,.62,'MEASURE & DECIDE  •  Randomize eligible accounts; review weekly; expand only with credible uplift and safe outcomes.',18,TEAL,True)
note(s,'Evaluation: define a consistent new-delinquency outcome and 90-day follow-up for every eligible account before assignment. Randomize at customer level within risk strata to support-versus-business-as-usual, preserving required notices and standard support access. Pre-specify sample size and a meaningful minimum effect; the supplied 500-row extract may be insufficient. Report absolute and relative changes, confidence intervals and differences in case mix; do not attribute recovery changes to AI from a before/after comparison alone. If the concurrent control has 16% delinquency, a 10% relative reduction means 14.4%, a 1.6 percentage-point reduction. The historical 16% is an illustration, not a measured future-incidence control baseline. Cost per resolved account = total collection operating cost, including agents and human reviews, divided by resolved accounts. Cure rate = delinquent accounts returned to current status divided by delinquent accounts enrolled, over a fixed follow-up; report separately from prevention. Net recovery deducts collection costs. Complaint rate uses contacts or customers as a consistent denominator; survey customer satisfaction and track opt-out fulfillment and support uptake. Monitor model discrimination, calibration, data drift and subgroup error rates. Collections owns outcomes, Finance verifies costs, Model Risk verifies prediction and fairness, Compliance owns contact incidents. Weekly controls review and a day-90 executive gate determine expand, revise or stop. No expansion if material customer harm, unresolved compliance incidents, fairness breaches or weak evidence of benefit remain. Expected qualitative benefits are aspirations to test. Task 3 document contains no completed recommendation or SMART target; all pilot targets and interventions here are proposed.')

prs.core_properties.title='Geldium | Responsible AI-Powered Collections'
prs.core_properties.subject='Executive strategy briefing — five-slide concept'
prs.core_properties.author='Geldium Collections Strategy'
prs.save(OUT/'Geldium_AI_Collections_Strategy.pptx')
print(OUT/'Geldium_AI_Collections_Strategy.pptx')
