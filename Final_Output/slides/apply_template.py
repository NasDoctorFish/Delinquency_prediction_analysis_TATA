import sys
sys.path.insert(0, '/private/tmp/geldium_pptx_lib')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pathlib import Path

root=Path(__file__).resolve().parent.parent
p=Presentation(root/'Presentation_Template.pptx')
source=Presentation(root/'deliverables/Geldium_AI_Collections_Strategy.pptx')
slides=list(p.slides)
# Reuse the template's optional slide for the second workflow slide.
ids=p.slides._sldIdLst
last=ids[-1]; ids.remove(last); ids.insert(2,last)
def set_text(shape,lines,size=18):
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    for i,line in enumerate(lines):
        q=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        q.text=line; q.font.name='Arial'; q.font.size=Pt(size)
        q.space_after=Pt(13)
def title(s,text): set_text(s.shapes.title,[text],28)
def body(s,lines): set_text(s.shapes[1],lines,18)
def box(s,x,y,w,h,lines,size=17):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(sh,lines,size)
    return sh
def notes(s,i): s.notes_slide.notes_text_frame.text=source.slides[i].notes_slide.notes_text_frame.text
def footer(s,text): box(s,.53,4.91,8.62,.32,[text],10)

set_text(slides[0].shapes[0],['AI-Powered Collections Strategy'],32)
set_text(slides[0].shapes[1],['Responsible, explainable automation for Geldium\nExecutive briefing | Proposed pilot'],20)
slides[0].notes_slide.notes_text_frame.text='Deck outline: customer-data-to-learning workflow; risk and action policy; agent autonomy versus human oversight; responsible AI guardrails; expected business and customer impact. Read all three project Word documents. Task 2 is a model plan without measured validation results; Task 3 still contains unfilled prompts. Interventions and targets are proposals.'

s=slides[1]; title(s,'How the System Works: A Controlled Loop')
body(s,[
'1. INPUTS — Payment history, balances, income and utilization; add verified account status and contact preferences.',
'2. DECISION — Interpretable logistic regression estimates risk; policy checks select eligible actions or human review.',
'3. ACTION — Send approved reminders; offer a support conversation; route hardship and disputes to specialists.',
'4. LEARNING — Track repayment, complaints and fairness versus a control group; humans approve model updates.'
]); notes(s,0); footer(s,'Proposed workflow | Sources: EDA Summary Report and Task 2 Model Plan')

s=slides[5]; title(s,'How the System Works: Risk to Support')
body(s,[
'Validate first: fix monthly payment encoding, standardize employment labels, preserve missingness and confirm data timing.',
'Low validated risk → routine reminders. Elevated risk → proactive support invitation and staff-led repayment options.',
'Missing data or uncertain scores → verify facts. Hardship, disputes or contact restrictions → pause outreach and review.',
'Validate and calibrate the Task 2 model before setting thresholds; match escalation volumes to staff capacity.'
]); notes(s,1); footer(s,'Exploratory signal: missing loan balance = 24.1% delinquency vs. 16.0% overall; investigate before acting.')

s=slides[2]; title(s,'Role of Agentic AI')
s.shapes[1].text_frame.clear()
t=s.shapes.add_table(5,2, Inches(.53), Inches(1.62), Inches(8.94), Inches(2.95)).table
left=['Autonomous','Validate inputs and score eligible accounts.','Select approved channel, timing and template.','Send reminders; log responses; create staff tasks.','Monitor drift, complaints and control breaches.']
right=['Human Oversight','Approve models, thresholds and contact rules.','Review hardship, disputes and uncertain scores.','Authorize plans, settlements and adverse actions.','Approve retraining; own audits and shutdown.']
for i in range(5):
    for j,values in enumerate([left,right]):
        c=t.cell(i,j); c.margin_left=Inches(.13); c.margin_top=Inches(.09)
        c.fill.solid(); c.fill.fore_color.rgb=RGBColor.from_string('0277BD' if i==0 else ('E8F2F8' if i%2 else 'FFFFFF'))
        set_text(c,[values[i]],16)
        for q in c.text_frame.paragraphs:
            q.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '004065'); q.font.bold=i==0; q.space_after=Pt(0)
notes(s,2); footer(s,'Exception → pause outreach → named specialist → logged decision and customer explanation')

s=slides[3]; title(s,'Responsible AI Guardrails')
body(s,[
'Fairness — Audit subgroup errors and access to support; investigate age, location, employment and missingness proxies.',
'Explainability — Give plain-language reasons, data correction and appeal routes; log facts, model version and overrides.',
'Compliance & privacy — Approve contact limits, opt-outs and dispute rules by market and creditor role; minimize data and access.',
'Accountability — Use verified facts and approved templates; human review for hardship and adverse actions; enable shutdown.'
]); notes(s,3); footer(s,'Sources: NIST AI RMF; CFPB Debt Collection Rule FAQs | Applicability and source URLs in notes')

s=slides[4]; title(s,'Expected Business Impact')
s.shapes[1].text_frame.clear()
for x,h,lines in [(.53,'Business KPIs',[
'• Target 10% lower new delinquency relative to concurrent control.',
'• Target 15% lower cost per resolved account, including AI and staff costs.',
'• Track cure rate, net recovery and routine workload.'
]),(5.1,'Customer Outcomes',[
'• Earlier access to affordable repayment support.',
'• Clear messages, preferred channels and simple opt-outs.',
'• Complaints no worse than control; monitor satisfaction and subgroup fairness.'
])]:
    heading=box(s,x,1.65,4.25,.4,[h],20)
    heading.text_frame.paragraphs[0].font.bold=True
    box(s,x,2.23,4.25,2.33,lines,17)
notes(s,4); footer(s,'Proposed 90-day randomized pilot | Targets are hypotheses, not forecasts; expand only with safe, credible uplift.')

p.core_properties.title='Geldium AI-Powered Collections Strategy'
out=root/'deliverables/Geldium_AI_Collections_Strategy_Template.pptx'
p.save(out)
check=Presentation(out)
assert len(check.slides)==6
for s in check.slides:
    assert s.has_notes_slide
    for sh in s.shapes:
        assert sh.left+sh.width<=check.slide_width and sh.top+sh.height<=check.slide_height
        if sh.has_text_frame:
            assert 'Prompts to try' not in sh.text and '[Feel free' not in sh.text
print(out)
